from __future__ import annotations

import shutil
import math
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, messagebox, ttk
from typing import Any, Callable

from openpyxl.utils import get_column_letter

from .config import AppPaths, SalaryConfig, SalaryRates
from .models import CalculationResult, ExpectedTotals, WorkbookDocument
from .workbook import (
    CacheManager,
    SalaryCalculator,
    TemplateRepository,
    WorkbookLoader,
    WorkbookValidator,
    discover_excel_files,
    find_name_occurrences,
    infer_expected_totals,
    normalize_text,
    save_document_edits,
)


BG = "#F4F7FA"
PANEL = "#FFFFFF"
NAVY = "#17365D"
BLUE = "#2E75B6"
RED = "#FCE4E4"
RED_TEXT = "#9C0006"
GRID = "#D6DEE8"
MUTED = "#5F6B7A"


class TemplateSetupCancelled(Exception):
    """用户在首次使用时未提供模板。"""


class ScrollableSpreadsheet(ttk.Frame):
    """使用 Entry 网格实现可横向滚动、单元格级标红的轻量预览。"""

    def __init__(self, master: tk.Misc, wheel_direction: int = 1):
        super().__init__(master)
        self.wheel_direction = 1 if wheel_direction >= 0 else -1
        self.canvas = tk.Canvas(self, bg=PANEL, highlightthickness=0)
        self.v_scroll = ttk.Scrollbar(self, orient="vertical", command=self.canvas.yview)
        self.h_scroll = ttk.Scrollbar(self, orient="horizontal", command=self.canvas.xview)
        self.canvas.configure(yscrollcommand=self.v_scroll.set, xscrollcommand=self.h_scroll.set)
        self.inner = tk.Frame(self.canvas, bg=PANEL)
        self.window = self.canvas.create_window((0, 0), window=self.inner, anchor="nw")
        self.canvas.grid(row=0, column=0, sticky="nsew")
        self.v_scroll.grid(row=0, column=1, sticky="ns")
        self.h_scroll.grid(row=1, column=0, sticky="ew")
        self.rowconfigure(0, weight=1)
        self.columnconfigure(0, weight=1)
        self.inner.bind("<Configure>", self._update_scroll_region)
        self._bind_scroll(self.canvas)
        self._bind_scroll(self.inner)
        self.variables: list[list[tk.StringVar]] = []
        self._values: list[list[Any]] = []
        self._editable = False

    def _update_scroll_region(self, _event: tk.Event) -> None:
        self.canvas.configure(scrollregion=self.canvas.bbox("all"))

    def _bind_scroll(self, widget: tk.Misc) -> None:
        widget.bind("<MouseWheel>", self._on_mousewheel)
        widget.bind("<Shift-MouseWheel>", self._on_shift_mousewheel)
        widget.bind("<Button-4>", self._on_mousewheel)
        widget.bind("<Button-5>", self._on_mousewheel)
        widget.bind("<Shift-Button-4>", self._on_shift_mousewheel)
        widget.bind("<Shift-Button-5>", self._on_shift_mousewheel)

    @staticmethod
    def _wheel_units(event: tk.Event) -> int:
        number = getattr(event, "num", None)
        if number == 4:
            return 1
        if number == 5:
            return -1
        delta = int(getattr(event, "delta", 0))
        if delta == 0:
            return 0
        return 1 if delta > 0 else -1

    def _on_mousewheel(self, event: tk.Event) -> str:
        if int(getattr(event, "state", 0)) & 0x0001:
            return self._on_shift_mousewheel(event)
        self.canvas.yview_scroll(
            self._wheel_units(event) * self.wheel_direction,
            "units",
        )
        return "break"

    def _on_shift_mousewheel(self, event: tk.Event) -> str:
        self.canvas.xview_scroll(
            self._wheel_units(event) * self.wheel_direction,
            "units",
        )
        return "break"

    def clear(self) -> None:
        for child in self.inner.winfo_children():
            child.destroy()
        self.variables = []
        self._values = []

    def load(
        self,
        values: list[list[Any]],
        issue_positions: set[tuple[int, int]] | None = None,
        editable: bool = False,
    ) -> None:
        self.clear()
        self._values = values
        self._editable = editable
        issue_positions = issue_positions or set()
        if not values:
            tk.Label(
                self.inner,
                text="此文件无法显示表格内容，请查看问题说明。",
                bg=PANEL,
                fg=MUTED,
                padx=24,
                pady=24,
            ).grid(row=0, column=0)
            return

        width = max((len(row) for row in values), default=1)
        corner = tk.Label(
            self.inner,
            text="",
            width=5,
            bg="#E8EEF5",
            relief="solid",
            bd=1,
        )
        corner.grid(row=0, column=0, sticky="nsew")
        self._bind_scroll(corner)
        for col in range(1, width + 1):
            column_label = tk.Label(
                self.inner,
                text=get_column_letter(col),
                width=14,
                bg="#E8EEF5",
                fg=NAVY,
                font=("Microsoft YaHei UI", 9, "bold"),
                relief="solid",
                bd=1,
            )
            column_label.grid(row=0, column=col, sticky="nsew")
            self._bind_scroll(column_label)

        for row_index, row in enumerate(values, start=1):
            row_label = tk.Label(
                self.inner,
                text=str(row_index),
                width=5,
                bg="#E8EEF5",
                fg=NAVY,
                relief="solid",
                bd=1,
            )
            row_label.grid(row=row_index, column=0, sticky="nsew")
            self._bind_scroll(row_label)
            variable_row: list[tk.StringVar] = []
            for col_index in range(1, width + 1):
                value = row[col_index - 1] if col_index <= len(row) else None
                variable = tk.StringVar(value="" if value is None else str(value))
                is_issue = (row_index, col_index) in issue_positions
                is_header = row_index == 1
                background = RED if is_issue else ("#DCE6F1" if is_header else PANEL)
                foreground = RED_TEXT if is_issue else NAVY if is_header else "#1F2937"
                entry = tk.Entry(
                    self.inner,
                    textvariable=variable,
                    width=14,
                    relief="solid",
                    bd=1,
                    bg=background,
                    fg=foreground,
                    readonlybackground=background,
                    disabledbackground=background,
                    disabledforeground=foreground,
                    font=("Microsoft YaHei UI", 9, "bold" if is_header else "normal"),
                    justify="center",
                )
                if not editable:
                    entry.configure(state="readonly")
                entry.grid(row=row_index, column=col_index, sticky="nsew")
                self._bind_scroll(entry)
                variable_row.append(variable)
            self.variables.append(variable_row)
        self.canvas.xview_moveto(0)
        self.canvas.yview_moveto(0)

    def values(self) -> list[list[Any]]:
        result: list[list[Any]] = []
        for variable_row in self.variables:
            row: list[Any] = []
            for variable in variable_row:
                value = variable.get()
                row.append(None if value == "" else value)
            result.append(row)
        return result


class SalaryDialog(tk.Toplevel):
    def __init__(
        self,
        master: tk.Misc,
        rates: SalaryRates,
        on_save: Callable[[SalaryRates], None],
    ):
        super().__init__(master)
        self.title("工资标准")
        self.resizable(False, False)
        self.transient(master)
        self.grab_set()
        self.on_save = on_save
        self.variables: dict[str, tk.StringVar] = {}
        panel = ttk.Frame(self, padding=20)
        panel.grid(sticky="nsew")
        values = [
            ("g1", "高一晚自修", rates.grade_evening["1"]),
            ("g2", "高二晚自修", rates.grade_evening["2"]),
            ("g3", "高三晚自修", rates.grade_evening["3"]),
            ("reading", "晚读", rates.evening_reading),
            ("noon", "高中中午", rates.noon),
            ("afternoon", "全程下午", rates.afternoon),
        ]
        for row, (key, label, value) in enumerate(values):
            ttk.Label(panel, text=f"{label}（元/节）").grid(row=row, column=0, sticky="w", pady=5)
            variable = tk.StringVar(value=_format_number(value))
            ttk.Entry(panel, textvariable=variable, width=16).grid(row=row, column=1, padx=(18, 0), pady=5)
            self.variables[key] = variable
        ttk.Label(panel, text="保存后会同步更新 salary.json 和模板金额列表头。", foreground=MUTED).grid(
            row=len(values), column=0, columnspan=2, sticky="w", pady=(12, 8)
        )
        buttons = ttk.Frame(panel)
        buttons.grid(row=len(values) + 1, column=0, columnspan=2, sticky="e")
        ttk.Button(buttons, text="取消", command=self.destroy).pack(side="right")
        ttk.Button(buttons, text="保存", style="Accent.TButton", command=self._save).pack(
            side="right", padx=(0, 8)
        )

    def _save(self) -> None:
        try:
            parsed = {key: float(variable.get().strip()) for key, variable in self.variables.items()}
            if any(not math.isfinite(value) or value < 0 for value in parsed.values()):
                raise ValueError("工资必须是有限的非负数字")
            rates = SalaryRates(
                grade_evening={"1": parsed["g1"], "2": parsed["g2"], "3": parsed["g3"]},
                evening_reading=parsed["reading"],
                noon=parsed["noon"],
                afternoon=parsed["afternoon"],
            )
            self.on_save(rates)
        except ValueError as exc:
            messagebox.showerror("工资标准无效", str(exc), parent=self)
            return
        self.destroy()


def _format_number(value: float) -> str:
    return str(int(value)) if float(value).is_integer() else str(value)


class SalaryApp(tk.Tk):
    def __init__(self, paths: AppPaths):
        super().__init__()
        self.withdraw()
        self.paths = paths
        self.salary_config = SalaryConfig(paths.salary)
        self.rates = self.salary_config.load()
        self.cache = CacheManager(paths.cache_root)
        self.loader = WorkbookLoader()
        self.documents: list[WorkbookDocument] = []
        self.current_document_index: int | None = None
        self.result: CalculationResult | None = None

        self.title("教师坐班工资统计")
        self.geometry("1280x820")
        self.minsize(980, 660)
        self.configure(bg=BG)
        self.protocol("WM_DELETE_WINDOW", self._close)
        self._configure_styles()
        self._build_ui()
        self.deiconify()
        self.update_idletasks()
        self.lift()

        self._ensure_template()
        self.validator = WorkbookValidator.from_template(paths.template)
        self._set_status("请选择包含班级统计表的文件夹。")

    def _ensure_template(self) -> None:
        if self.paths.template.exists():
            return
        while not self.paths.template.exists():
            if not self._confirm_initial_template_selection():
                self.destroy()
                raise TemplateSetupCancelled("首次使用未选择工资模板")
            selected = filedialog.askopenfilename(
                parent=self,
                title="首次使用：请选择教师工资模板",
                filetypes=[("Excel 工作簿", "*.xlsx")],
            )
            if not selected:
                continue
            try:
                TemplateRepository.install_template(Path(selected), self.paths.template)
            except Exception as exc:
                messagebox.showerror(
                    "模板无效",
                    f"无法使用所选文件：\n{exc}\n\n请重新选择完整的 .xlsx 工资模板。",
                    parent=self,
                )
                continue
            messagebox.showinfo(
                "模板导入成功",
                "工资模板已复制到 docs/module.xlsx。",
                parent=self,
            )

    def _confirm_initial_template_selection(self) -> bool:
        result = {"choose": False}
        dialog = tk.Toplevel(self)
        dialog.title("缺少工资模板")
        dialog.transient(self)
        dialog.resizable(False, False)
        dialog.configure(bg=PANEL)

        panel = ttk.Frame(dialog, style="Panel.TFrame", padding=(26, 22))
        panel.pack(fill="both", expand=True)
        ttk.Label(
            panel,
            text="缺少工资模板",
            style="Panel.TLabel",
            font=("Microsoft YaHei UI", 13, "bold"),
        ).pack(anchor="w")
        ttk.Label(
            panel,
            text=(
                "这是第一次使用本程序，需要先提供一个完整的 .xlsx 工资模板。\n"
                "程序会校验模板，并将其复制为 docs/module.xlsx。"
            ),
            style="Panel.TLabel",
            justify="left",
        ).pack(anchor="w", pady=(10, 20))

        buttons = ttk.Frame(panel, style="Panel.TFrame")
        buttons.pack(fill="x")

        def finish(choose: bool) -> None:
            result["choose"] = choose
            dialog.destroy()

        ttk.Button(buttons, text="退出程序", command=lambda: finish(False)).pack(side="right")
        choose_button = ttk.Button(
            buttons,
            text="立即选择模板",
            style="Accent.TButton",
            command=lambda: finish(True),
        )
        choose_button.pack(side="right", padx=(0, 8))
        dialog.protocol("WM_DELETE_WINDOW", lambda: finish(False))
        dialog.bind("<Return>", lambda _event: finish(True))
        dialog.bind("<Escape>", lambda _event: finish(False))

        dialog.update_idletasks()
        x = self.winfo_rootx() + max(0, (self.winfo_width() - dialog.winfo_reqwidth()) // 2)
        y = self.winfo_rooty() + max(0, (self.winfo_height() - dialog.winfo_reqheight()) // 2)
        dialog.geometry(f"+{x}+{y}")
        dialog.grab_set()
        choose_button.focus_set()
        self.wait_window(dialog)
        return result["choose"]

    def _configure_styles(self) -> None:
        style = ttk.Style(self)
        if "vista" in style.theme_names():
            style.theme_use("vista")
        style.configure("TFrame", background=BG)
        style.configure("Panel.TFrame", background=PANEL)
        style.configure("TLabel", background=BG, font=("Microsoft YaHei UI", 9))
        style.configure("Panel.TLabel", background=PANEL, font=("Microsoft YaHei UI", 9))
        style.configure("Title.TLabel", background=NAVY, foreground="white", font=("Microsoft YaHei UI", 18, "bold"))
        style.configure("Subtitle.TLabel", background=NAVY, foreground="#DCE6F1", font=("Microsoft YaHei UI", 9))
        style.configure("Accent.TButton", font=("Microsoft YaHei UI", 9, "bold"))
        style.configure("Treeview", rowheight=28, font=("Microsoft YaHei UI", 9))
        style.configure("Treeview.Heading", font=("Microsoft YaHei UI", 9, "bold"))

    def _build_ui(self) -> None:
        self.columnconfigure(0, weight=1)
        self.rowconfigure(2, weight=1)

        header = ttk.Frame(self, style="Panel.TFrame")
        header.configure(height=86)
        header.grid(row=0, column=0, sticky="ew")
        banner = tk.Frame(header, bg=NAVY, height=86)
        banner.pack(fill="both", expand=True)
        ttk.Label(banner, text="教师坐班工资统计", style="Title.TLabel").pack(anchor="w", padx=24, pady=(14, 2))
        ttk.Label(
            banner,
            text="批量校验班级表，修正问题后生成并预览工资统计表",
            style="Subtitle.TLabel",
        ).pack(anchor="w", padx=25)

        controls = ttk.Frame(self, style="Panel.TFrame", padding=(18, 14))
        controls.grid(row=1, column=0, sticky="ew", padx=16, pady=(14, 8))
        controls.columnconfigure(1, weight=1)
        ttk.Label(controls, text="导入文件夹", style="Panel.TLabel").grid(row=0, column=0, sticky="w")
        self.folder_var = tk.StringVar()
        ttk.Entry(controls, textvariable=self.folder_var, state="readonly").grid(
            row=0, column=1, sticky="ew", padx=10
        )
        ttk.Button(controls, text="选择文件夹", command=self._choose_folder).grid(row=0, column=2, padx=(0, 8))
        ttk.Button(controls, text="工资标准", command=self._open_salary_dialog).grid(
            row=0, column=3, padx=(0, 8)
        )
        ttk.Button(controls, text="更换模板", command=self._replace_template).grid(row=0, column=4)

        options = ttk.Frame(controls, style="Panel.TFrame")
        options.grid(row=1, column=0, columnspan=5, sticky="ew", pady=(12, 0))
        ttk.Label(options, text="年级", style="Panel.TLabel").pack(side="left")
        self.grade_var = tk.StringVar(value="1")
        ttk.Combobox(options, textvariable=self.grade_var, values=("1", "2", "3"), state="readonly", width=5).pack(
            side="left", padx=(6, 22)
        )
        self.expected_vars = {
            "evening": tk.StringVar(value="5"),
            "noon": tk.StringVar(value="5"),
            "afternoon": tk.StringVar(value="5"),
        }
        for text, key in (("晚自修总节数", "evening"), ("中午总节数", "noon"), ("下午总节数", "afternoon")):
            ttk.Label(options, text=text, style="Panel.TLabel").pack(side="left")
            ttk.Entry(options, textvariable=self.expected_vars[key], width=7).pack(side="left", padx=(6, 18))
        ttk.Button(
            options,
            text="校验并生成预览",
            style="Accent.TButton",
            command=self._calculate,
        ).pack(side="right")

        self.notebook = ttk.Notebook(self)
        self.notebook.grid(row=2, column=0, sticky="nsew", padx=16, pady=8)
        self.issue_tab = ttk.Frame(self.notebook, padding=10)
        self.output_tab = ttk.Frame(self.notebook, padding=10)
        self.notebook.add(self.issue_tab, text="导入与问题")
        self.notebook.add(self.output_tab, text="工资预览")
        self._build_issue_tab()
        self._build_output_tab()

        self.footer = ttk.Frame(self, padding=(16, 4, 16, 12))
        self.footer.grid(row=3, column=0, sticky="ew")
        self.footer.columnconfigure(0, weight=1)
        self.status_var = tk.StringVar()
        ttk.Label(
            self.footer,
            textvariable=self.status_var,
            foreground=MUTED,
            anchor="w",
        ).grid(row=0, column=0, sticky="ew", padx=(0, 12))
        self.save_button = ttk.Button(
            self.footer,
            text="导出 Excel…",
            command=self._save_as,
            state="disabled",
        )
        self.save_button.grid(row=0, column=1, sticky="e")

    def _build_issue_tab(self) -> None:
        self.issue_tab.columnconfigure(1, weight=1)
        self.issue_tab.rowconfigure(0, weight=1)
        left = ttk.Frame(self.issue_tab, style="Panel.TFrame", padding=8)
        left.grid(row=0, column=0, sticky="nsw", padx=(0, 10))
        ttk.Label(left, text="导入文件", style="Panel.TLabel", font=("Microsoft YaHei UI", 10, "bold")).pack(
            anchor="w", pady=(0, 8)
        )
        self.file_tree = ttk.Treeview(left, columns=("status", "count"), show="tree headings", height=18)
        self.file_tree.heading("#0", text="文件")
        self.file_tree.heading("status", text="状态")
        self.file_tree.heading("count", text="问题")
        self.file_tree.column("#0", width=190)
        self.file_tree.column("status", width=70, anchor="center")
        self.file_tree.column("count", width=55, anchor="center")
        self.file_tree.pack(fill="y", expand=True)
        self.file_tree.bind("<<TreeviewSelect>>", self._select_document)

        right = ttk.Frame(self.issue_tab, style="Panel.TFrame", padding=8)
        right.grid(row=0, column=1, sticky="nsew")
        right.columnconfigure(0, weight=1)
        right.rowconfigure(1, weight=1)
        self.current_file_var = tk.StringVar(value="尚未导入文件")
        ttk.Label(right, textvariable=self.current_file_var, style="Panel.TLabel", font=("Microsoft YaHei UI", 10, "bold")).grid(
            row=0, column=0, sticky="w", pady=(0, 8)
        )
        self.input_grid = ScrollableSpreadsheet(right)
        self.input_grid.grid(row=1, column=0, sticky="nsew")
        issue_panel = ttk.Frame(right, style="Panel.TFrame")
        issue_panel.grid(row=2, column=0, sticky="ew", pady=(8, 0))
        issue_panel.columnconfigure(0, weight=1)
        self.issue_text = tk.Text(
            issue_panel,
            height=4,
            wrap="word",
            relief="solid",
            bd=1,
            bg="#FFF8F0",
            fg="#7F3B08",
            font=("Microsoft YaHei UI", 9),
        )
        self.issue_text.grid(row=0, column=0, sticky="ew", padx=(0, 8))
        self.issue_text.configure(state="disabled")
        self.recommend_button = ttk.Button(
            issue_panel,
            text="使用推荐姓名",
            command=self._apply_recommended_name,
            state="disabled",
        )
        self.recommend_button.grid(row=0, column=1, sticky="ns", padx=(0, 8))
        self.append_template_button = ttk.Button(
            issue_panel,
            text="追加到模板",
            command=self._append_problem_name_to_template,
            state="disabled",
        )
        self.append_template_button.grid(row=0, column=2, sticky="ns", padx=(0, 8))
        ttk.Button(issue_panel, text="应用编辑并重新校验", command=self._apply_edits).grid(
            row=0, column=3, sticky="ns"
        )

    def _build_output_tab(self) -> None:
        self.output_tab.columnconfigure(0, weight=1)
        self.output_tab.rowconfigure(1, weight=1)
        self.output_summary_var = tk.StringVar(value="校验通过后将在此显示工资表。")
        ttk.Label(
            self.output_tab,
            textvariable=self.output_summary_var,
            font=("Microsoft YaHei UI", 10, "bold"),
        ).grid(row=0, column=0, sticky="w", pady=(0, 8))
        # Windows 下只读 Entry 的滚轮方向与可编辑表格表现相反。
        self.output_grid = ScrollableSpreadsheet(self.output_tab, wheel_direction=-1)
        self.output_grid.grid(row=1, column=0, sticky="nsew")

    def _choose_folder(self) -> None:
        initial = self.folder_var.get() or str(self.paths.root / "example")
        selected = filedialog.askdirectory(title="选择班级统计表文件夹", initialdir=initial)
        if not selected:
            return
        self._sync_current_document(save=False)
        folder = Path(selected)
        files = discover_excel_files(folder)
        if not files:
            messagebox.showwarning("没有 Excel 文件", "所选文件夹中没有 .xlsx、.xlsm 或 .xls 文件。")
            return
        self.cache.cleanup()
        self.documents = []
        self.current_document_index = None
        self.result = None
        self.save_button.configure(state="disabled")
        self.output_grid.clear()
        self.output_summary_var.set("校验通过后将在此显示工资表。")
        self.folder_var.set(str(folder))
        for index, path in enumerate(files, start=1):
            cache_path = self.cache.copy_in(path, prefix=f"{index:03d}-")
            document = self.loader.load(path, cache_path)
            self.documents.append(document)
        fallback = self._expected_totals(show_error=False) or ExpectedTotals(5, 5, 5)
        expected = infer_expected_totals(self.documents, fallback)
        self._set_expected_totals(expected)
        for document in self.documents:
            self.validator.validate(document, expected)
        self._refresh_file_tree()
        self._set_status(
            f"{self._status_summary()} 默认核对值已按班级实际合计的众数设置："
            f"晚自修 {_format_number(expected.evening)}，"
            f"中午 {_format_number(expected.noon)}，"
            f"下午 {_format_number(expected.afternoon)}。"
        )
        if self.documents:
            first_issue = next((i for i, doc in enumerate(self.documents) if doc.issues), 0)
            self.file_tree.selection_set(str(first_issue))
            self.file_tree.focus(str(first_issue))
            self._show_document(first_issue)

    def _expected_totals(self, show_error: bool) -> ExpectedTotals | None:
        try:
            values = {key: float(variable.get().strip()) for key, variable in self.expected_vars.items()}
            if any(not math.isfinite(value) or value < 0 for value in values.values()):
                raise ValueError("核对总节数必须是有限的非负数字")
            return ExpectedTotals(**values)
        except ValueError as exc:
            if show_error:
                messagebox.showerror("核对条件无效", f"三个核对总节数都必须是非负数字。\n{exc}")
            return None

    def _set_expected_totals(self, expected: ExpectedTotals) -> None:
        self.expected_vars["evening"].set(_format_number(expected.evening))
        self.expected_vars["noon"].set(_format_number(expected.noon))
        self.expected_vars["afternoon"].set(_format_number(expected.afternoon))

    def _refresh_file_tree(self) -> None:
        selected = self.current_document_index
        for item in self.file_tree.get_children():
            self.file_tree.delete(item)
        for index, document in enumerate(self.documents):
            self.file_tree.insert(
                "",
                "end",
                iid=str(index),
                text=document.filename,
                values=("需处理" if document.issues else "通过", len(document.issues)),
            )
        if selected is not None and selected < len(self.documents):
            self.file_tree.selection_set(str(selected))

    def _select_document(self, _event: tk.Event) -> None:
        selection = self.file_tree.selection()
        if not selection:
            return
        next_index = int(selection[0])
        if self.current_document_index is not None and self.current_document_index != next_index:
            if not self._try_sync_current_document(save=True):
                self.file_tree.selection_set(str(self.current_document_index))
                return
        self._show_document(next_index)

    def _show_document(self, index: int) -> None:
        if not (0 <= index < len(self.documents)):
            return
        self.current_document_index = index
        document = self.documents[index]
        self.current_file_var.set(
            f"{document.filename}    工作表：{document.sheet_name or '无法读取'}"
        )
        self.input_grid.load(document.values, document.issue_positions, editable=not bool(document.load_error))
        self._set_issue_text(document)

    def _set_issue_text(self, document: WorkbookDocument) -> None:
        if document.issues:
            lines = []
            for issue in document.issues:
                position = (
                    f"{get_column_letter(issue.column)}{issue.row}"
                    if issue.row > 0 and issue.column > 0
                    else "文件"
                )
                lines.append(f"{position}：{issue.message}")
            content = "\n".join(lines)
        else:
            content = "校验通过。"
        self.issue_text.configure(state="normal")
        self.issue_text.delete("1.0", "end")
        self.issue_text.insert("1.0", content)
        self.issue_text.configure(state="disabled")
        recommendation = self._first_name_recommendation(document)
        if recommendation:
            _, _, suggested_name, similarity = recommendation
            self.recommend_button.configure(
                text=f"替换为“{suggested_name}” ({similarity:.0%})",
                state="normal",
            )
        else:
            self.recommend_button.configure(text="使用推荐姓名", state="disabled")
        problem_name = self._first_problem_name(document)
        if problem_name:
            _, _, name = problem_name
            self.append_template_button.configure(
                text=f"追加“{name}”到模板",
                state="normal",
            )
        else:
            self.append_template_button.configure(text="追加到模板", state="disabled")

    @staticmethod
    def _first_problem_name(document: WorkbookDocument) -> tuple[int, int, str] | None:
        for issue in document.issues:
            if issue.kind != "name" or issue.row <= 0 or issue.column <= 0:
                continue
            name = normalize_text(document.cell(issue.row, issue.column))
            if name:
                return issue.row, issue.column, name
        return None

    def _first_name_recommendation(
        self, document: WorkbookDocument
    ) -> tuple[int, int, str, float] | None:
        for issue in document.issues:
            if issue.kind != "name" or issue.row <= 0 or issue.column <= 0:
                continue
            recommendation = self.validator.recommend_name(document.cell(issue.row, issue.column))
            if recommendation:
                suggested_name, similarity = recommendation
                return issue.row, issue.column, suggested_name, similarity
        return None

    def _apply_recommended_name(self) -> None:
        if self.current_document_index is None:
            return
        document = self.documents[self.current_document_index]
        recommendation = self._first_name_recommendation(document)
        if not recommendation:
            messagebox.showinfo("没有推荐姓名", "当前文件没有可直接替换的姓名推荐。")
            return
        row, column, suggested_name, _ = recommendation
        try:
            self.input_grid.variables[row - 1][column - 1].set(suggested_name)
        except IndexError:
            messagebox.showerror("替换失败", "无法定位需要替换的姓名单元格。")
            return
        self._apply_edits()

    def _append_problem_name_to_template(self) -> None:
        if self.current_document_index is None:
            return
        expected = self._expected_totals(show_error=True)
        if expected is None or not self._try_sync_current_document(save=True):
            return
        for document in self.documents:
            self.validator.validate(document, expected)
        current = self.documents[self.current_document_index]
        problem = self._first_problem_name(current)
        if not problem:
            self._refresh_file_tree()
            self._show_document(self.current_document_index)
            messagebox.showinfo("无需追加", "当前文件没有可追加到模板的姓名问题。")
            return
        _, _, name = problem
        occurrences = find_name_occurrences(self.documents, name)
        confirmed = messagebox.askyesno(
            "追加人员到模板",
            f"是否将“{name}”永久追加到 docs/module.xlsx 的末尾？\n\n"
            f"添加后，当前导入文件中的 {len(occurrences)} 个同名记录都会视为有效。",
            parent=self,
        )
        if not confirmed:
            return
        try:
            new_row = TemplateRepository(self.paths.template).append_person(name)
            self.validator = WorkbookValidator.from_template(self.paths.template)
            for document in self.documents:
                self.validator.validate(document, expected)
        except Exception as exc:
            self._set_status(f"追加模板人员失败：{exc}")
            messagebox.showerror(
                "追加失败",
                f"无法将“{name}”追加到模板：\n{exc}",
                parent=self,
            )
            return
        self._refresh_file_tree()
        self._show_document(self.current_document_index)
        self._invalidate_result()
        self._set_status(
            f"已将“{name}”追加到模板第 {new_row} 行；{len(occurrences)} 个同名记录已通过姓名校验。"
        )
        messagebox.showinfo(
            "追加成功",
            f"已将“{name}”追加到模板第 {new_row} 行。\n"
            f"{len(occurrences)} 个同名记录已重新校验。",
            parent=self,
        )

    def _sync_current_document(self, save: bool) -> None:
        if self.current_document_index is None or not self.input_grid.variables:
            return
        document = self.documents[self.current_document_index]
        if document.load_error:
            return
        document.values = self.input_grid.values()
        if save:
            save_document_edits(document)

    def _try_sync_current_document(self, save: bool) -> bool:
        try:
            self._sync_current_document(save=save)
            return True
        except Exception as exc:
            self._set_status(f"保存编辑失败：{exc}")
            messagebox.showerror("保存编辑失败", f"无法保存当前修改：\n{exc}")
            return False

    def _apply_edits(self) -> None:
        if self.current_document_index is None:
            return
        expected = self._expected_totals(show_error=True)
        if expected is None:
            return
        corrections = self._name_corrections_from_current_grid()
        if not self._try_sync_current_document(save=True):
            return
        batch_count = self._offer_batch_name_replacement(corrections)
        for document in self.documents:
            self.validator.validate(document, expected)
        self._refresh_file_tree()
        self._show_document(self.current_document_index)
        self._invalidate_result()
        self._set_status(self._status_summary())
        current = self.documents[self.current_document_index]
        batch_message = f"，并批量修改了 {batch_count} 个同名单元格" if batch_count else ""
        if current.issues:
            messagebox.showwarning(
                "重新校验完成",
                f"修改已保存{batch_message}。当前文件仍有 {len(current.issues)} 个问题。\n"
                "请继续处理红色单元格。",
            )
        else:
            messagebox.showinfo(
                "重新校验完成",
                f"修改已保存{batch_message}，当前文件已通过校验。",
            )

    def _name_corrections_from_current_grid(self) -> list[tuple[str, str]]:
        """提取本次把错误姓名改成有效模板姓名的操作。"""
        if self.current_document_index is None:
            return []
        document = self.documents[self.current_document_index]
        edited_values = self.input_grid.values()
        corrections: dict[tuple[str, str], str] = {}
        for issue in document.issues:
            if issue.kind != "name" or issue.row <= 0 or issue.column <= 0:
                continue
            old_name = normalize_text(document.cell(issue.row, issue.column))
            try:
                raw_new_name = edited_values[issue.row - 1][issue.column - 1]
            except IndexError:
                continue
            new_name = normalize_text(raw_new_name)
            if (
                old_name
                and new_name
                and old_name != new_name
                and new_name in self.validator.template_names
            ):
                display_name = str(raw_new_name).strip()
                corrections[(old_name, new_name)] = display_name
        return [(old_name, display_name) for (old_name, _), display_name in corrections.items()]

    def _offer_batch_name_replacement(self, corrections: list[tuple[str, str]]) -> int:
        replaced = 0
        for old_name, new_name in corrections:
            matches = find_name_occurrences(self.documents, old_name)
            if not matches:
                continue
            affected_files = len({document.source_path for document, _, _ in matches})
            confirmed = messagebox.askyesno(
                "批量修改同名错误",
                f"还发现 {len(matches)} 个“{old_name}”，分布在 {affected_files} 个文件中。\n\n"
                f"是否全部修改为“{new_name}”？",
                parent=self,
            )
            if not confirmed:
                continue
            changed_documents: dict[Path, WorkbookDocument] = {}
            for document, row, column in matches:
                document.set_cell(row, column, new_name)
                changed_documents[document.source_path] = document
            try:
                for document in changed_documents.values():
                    save_document_edits(document)
            except Exception as exc:
                self._set_status(f"批量保存姓名失败：{exc}")
                messagebox.showerror(
                    "批量修改失败",
                    f"无法保存部分批量修改：\n{exc}",
                    parent=self,
                )
                continue
            replaced += len(matches)
        return replaced

    def _calculate(self) -> None:
        if not self.documents:
            messagebox.showwarning("尚未导入", "请先选择包含班级统计表的文件夹。")
            return
        expected = self._expected_totals(show_error=True)
        if expected is None:
            return
        if not self._try_sync_current_document(save=True):
            return
        for document in self.documents:
            self.validator.validate(document, expected)
        self._refresh_file_tree()
        issue_count = sum(len(document.issues) for document in self.documents)
        if issue_count:
            first_issue = next(i for i, doc in enumerate(self.documents) if doc.issues)
            self.file_tree.selection_set(str(first_issue))
            self._show_document(first_issue)
            self.notebook.select(self.issue_tab)
            self._invalidate_result()
            messagebox.showwarning(
                "发现数据问题",
                f"共发现 {issue_count} 个问题。请根据红色单元格和问题说明修正后重新生成。",
            )
            self._set_status(self._status_summary())
            return
        try:
            output_path = self.cache.output_path()
            calculator = SalaryCalculator(self.paths.template, self.rates)
            self.result = calculator.calculate(
                self.documents,
                self.grade_var.get(),
                output_path,
            )
        except Exception as exc:
            self._invalidate_result()
            messagebox.showerror("生成失败", str(exc))
            return
        self.output_grid.load(self.result.values, editable=False)
        self.output_summary_var.set(
            f"已汇总 {self.result.source_file_count} 个班级文件，"
            f"涉及 {self.result.teacher_count} 位教师，合计金额 {_format_number(self.result.total_amount)} 元。"
        )
        self.save_button.configure(state="normal")
        self.notebook.select(self.output_tab)
        self._set_status("工资表已生成到临时缓存。确认预览后可点击“导出 Excel…”。")

    def _invalidate_result(self) -> None:
        self.result = None
        self.save_button.configure(state="disabled")
        self.output_grid.clear()
        self.output_summary_var.set("数据或设置已变化，请重新生成工资预览。")

    def _save_as(self) -> None:
        if self.result is None or not self.result.workbook_path.exists():
            messagebox.showwarning("没有可保存内容", "请先生成工资预览。")
            return
        destination = filedialog.asksaveasfilename(
            title="导出工资统计表",
            defaultextension=".xlsx",
            filetypes=[("Excel 工作簿", "*.xlsx")],
            initialfile="教师津贴统计表.xlsx",
        )
        if not destination:
            return
        try:
            shutil.copy2(self.result.workbook_path, destination)
        except OSError as exc:
            messagebox.showerror("保存失败", str(exc))
            return
        self._set_status(f"已保存：{destination}")
        messagebox.showinfo("导出成功", f"工资统计表已导出到：\n{destination}")

    def _open_salary_dialog(self) -> None:
        SalaryDialog(self, self.rates, self._save_rates)

    def _replace_template(self) -> None:
        confirmed = messagebox.askyesno(
            "更换工资模板",
            "更换模板后，姓名名单和工资表结构将以新模板为准。\n\n是否立即选择新模板？",
            parent=self,
        )
        if not confirmed:
            return
        selected = filedialog.askopenfilename(
            parent=self,
            title="选择新的教师工资模板",
            initialdir=str(self.paths.template.parent),
            filetypes=[("Excel 工作簿", "*.xlsx")],
        )
        if not selected:
            return
        if not self._try_sync_current_document(save=True):
            return
        try:
            TemplateRepository.replace_template(
                Path(selected),
                self.paths.template,
                self.rates,
            )
            self.validator = WorkbookValidator.from_template(self.paths.template)
            expected = self._expected_totals(show_error=False) or ExpectedTotals(5, 5, 5)
            for document in self.documents:
                self.validator.validate(document, expected)
        except Exception as exc:
            messagebox.showerror(
                "更换模板失败",
                f"无法使用所选模板：\n{exc}\n\n原模板未被候选文件直接覆盖。",
                parent=self,
            )
            return
        self._refresh_file_tree()
        if self.current_document_index is not None:
            self._show_document(self.current_document_index)
        self._invalidate_result()
        self._set_status("工资模板已更换，当前导入文件已按新模板重新校验。")
        messagebox.showinfo(
            "模板已更换",
            "新模板已保存为 docs/module.xlsx，并已同步当前工资标准。",
            parent=self,
        )

    def _save_rates(self, rates: SalaryRates) -> None:
        try:
            TemplateRepository(self.paths.template).update_rate_headers(rates)
            self.salary_config.save(rates)
        except Exception as exc:
            raise ValueError(f"保存失败：{exc}") from exc
        self.rates = rates
        self.validator = WorkbookValidator.from_template(self.paths.template)
        self._invalidate_result()
        self._set_status("工资标准已保存，并同步更新模板表头。")

    def _status_summary(self) -> str:
        files = len(self.documents)
        issues = sum(len(document.issues) for document in self.documents)
        affected = sum(1 for document in self.documents if document.issues)
        if issues:
            return f"已导入 {files} 个文件；{affected} 个文件共 {issues} 个问题待处理。"
        return f"已导入并通过校验 {files} 个文件。"

    def _set_status(self, text: str) -> None:
        self.status_var.set(text)

    def _close(self) -> None:
        try:
            self.cache.cleanup()
        finally:
            self.destroy()
