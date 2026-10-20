from __future__ import annotations

import math
import os
import re
import shutil
import stat
import uuid
from collections import Counter, defaultdict
from copy import copy
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .config import SalaryRates
from .models import CalculationResult, CellIssue, ExpectedTotals, WorkbookDocument


REQUIRED_COLUMNS = ("姓名", "晚自修", "高中中午", "全程下午")
WEEKEND_MONITOR_COLUMN = "周六日监考"
WEEKEND_COLUMN = "周六日"
HOLIDAY_COLUMN = "节假日"
WEEKEND_SOURCE_COLUMN = "周六日上课"
HOLIDAY_SOURCE_COLUMN = "节假日补课"
SUPPLEMENTAL_COLUMNS = ("姓名", WEEKEND_COLUMN, HOLIDAY_COLUMN)
CLASS_DOCUMENT = "class"
SUPPLEMENTAL_DOCUMENT = "weekend_holiday"
SUPPORTED_SUFFIXES = {".xlsx", ".xlsm"}
ALL_EXCEL_SUFFIXES = {".xlsx", ".xlsm", ".xls"}


def make_writable(path: Path) -> None:
    """清除复制文件继承的只读属性，尤其是 Windows 上的 Excel 文件。"""
    path.chmod(path.stat().st_mode | stat.S_IWRITE)


def normalize_text(value: Any) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", "", str(value)).strip()


def _is_weekend_source_header(label: str) -> bool:
    return label in {WEEKEND_COLUMN, WEEKEND_SOURCE_COLUMN} or (
        "周六" in label and "上课" in label
    )


def _is_holiday_source_header(label: str) -> bool:
    return label in {HOLIDAY_COLUMN, HOLIDAY_SOURCE_COLUMN} or "节假日" in label


def recommend_template_name(
    name: Any, template_names: Iterable[str]
) -> tuple[str, float] | None:
    """返回一个模板姓名推荐及 0~1 的相似度评分。"""
    target = normalize_text(name)
    candidates = sorted({normalize_text(item) for item in template_names if normalize_text(item)})
    if not target or not candidates:
        return None

    def score(candidate: str) -> float:
        text_similarity = SequenceMatcher(None, target, candidate).ratio()
        same_surname = 1.0 if target[0] == candidate[0] else 0.0
        same_length = 1.0 if len(target) == len(candidate) else 0.0
        return 0.6 * text_similarity + 0.25 * same_surname + 0.15 * same_length

    recommended = max(candidates, key=lambda candidate: (score(candidate), candidate))
    return recommended, score(recommended)


def number_from_cell(value: Any) -> float:
    if value is None or (isinstance(value, str) and not value.strip()):
        return 0.0
    if isinstance(value, bool):
        raise ValueError("布尔值不是有效节数")
    if isinstance(value, (int, float)):
        result = float(value)
    else:
        text = str(value).strip().replace(",", "")
        result = float(text)
    if not math.isfinite(result):
        raise ValueError("数值必须是有限数字")
    if result < 0:
        raise ValueError("节数不能小于 0")
    return result


def display_number(value: float) -> int | float:
    return int(value) if float(value).is_integer() else round(float(value), 4)


def document_metric_totals(doc: WorkbookDocument) -> dict[str, float]:
    """读取单个班级文件的实际合计；无效列不会参与默认值推断。"""
    if (
        doc.document_type != CLASS_DOCUMENT
        or doc.load_error
        or any(label not in doc.columns for label in REQUIRED_COLUMNS)
    ):
        return {}
    has_data = False
    totals: dict[str, float] = {}
    metric_labels = list(REQUIRED_COLUMNS[1:])
    if WEEKEND_MONITOR_COLUMN in doc.columns:
        metric_labels.append(WEEKEND_MONITOR_COLUMN)
    for label in metric_labels:
        total = 0.0
        valid = True
        column = doc.columns[label]
        for row in range(doc.header_row + 1, len(doc.values) + 1):
            name_value = doc.cell(row, doc.columns["姓名"])
            metric_values = [doc.cell(row, doc.columns[item]) for item in metric_labels]
            row_has_data = bool(normalize_text(name_value)) or any(
                value is not None and not (isinstance(value, str) and not value.strip())
                for value in metric_values
            )
            if not row_has_data:
                continue
            has_data = True
            try:
                total += number_from_cell(doc.cell(row, column))
            except (TypeError, ValueError):
                valid = False
                break
        if valid and has_data:
            totals[label] = total
    return totals


def infer_expected_totals(
    documents: Iterable[WorkbookDocument],
    fallback: ExpectedTotals | None = None,
) -> ExpectedTotals:
    """按各班实际合计的众数推断四个默认核对值。并列时采用先导入的值。"""
    fallback = fallback or ExpectedTotals(0, 0, 0)
    metric_labels = (*REQUIRED_COLUMNS[1:], WEEKEND_MONITOR_COLUMN)
    values: dict[str, list[float]] = {label: [] for label in metric_labels}
    for document in documents:
        for label, total in document_metric_totals(document).items():
            values[label].append(round(total, 9))

    def most_common(label: str, default: float) -> float:
        candidates = values[label]
        if not candidates:
            return default
        counts = Counter(candidates)
        highest = max(counts.values())
        return next(value for value in candidates if counts[value] == highest)

    return ExpectedTotals(
        evening=most_common("晚自修", fallback.evening),
        noon=most_common("高中中午", fallback.noon),
        afternoon=most_common("全程下午", fallback.afternoon),
        weekend_monitor=most_common(
            WEEKEND_MONITOR_COLUMN, fallback.weekend_monitor
        ),
    )


class CacheManager:
    """只管理本次运行创建的缓存目录，避免误删其他内容。"""

    def __init__(self, cache_root: Path):
        self.cache_root = cache_root
        self.session_dir: Path | None = None

    def start(self) -> Path:
        if self.session_dir is None:
            self.cache_root.mkdir(parents=True, exist_ok=True)
            self.session_dir = self.cache_root / f"session-{uuid.uuid4().hex}"
            self.session_dir.mkdir(parents=True)
        return self.session_dir

    def copy_in(self, source: Path, prefix: str = "") -> Path:
        session = self.start()
        destination = session / f"{prefix}{source.name}"
        counter = 1
        while destination.exists():
            destination = session / f"{prefix}{source.stem}-{counter}{source.suffix}"
            counter += 1
        shutil.copy2(source, destination)
        make_writable(destination)
        return destination

    def output_path(self, filename: str = "教师津贴统计表.xlsx") -> Path:
        return self.start() / filename

    def cleanup(self) -> None:
        if self.session_dir and self.session_dir.exists():
            shutil.rmtree(self.session_dir)
        self.session_dir = None
        if self.cache_root.exists() and not any(self.cache_root.iterdir()):
            self.cache_root.rmdir()


class WorkbookLoader:
    def load(self, source_path: Path, cache_path: Path | None = None) -> WorkbookDocument:
        doc = WorkbookDocument(source_path=source_path, cache_path=cache_path)
        path = cache_path or source_path
        if source_path.suffix.lower() not in SUPPORTED_SUFFIXES:
            doc.load_error = "不支持旧版 .xls，请另存为 .xlsx 后重试"
            doc.issues = [CellIssue(0, 0, doc.load_error, "format")]
            return doc
        try:
            workbook = load_workbook(path, data_only=True, read_only=True)
            worksheet, header_row, columns, document_type = self._find_data_sheet(workbook)
            doc.sheet_name = worksheet.title
            doc.header_row = header_row
            doc.columns = columns
            doc.document_type = document_type
            last_row, last_col = self._used_bounds(worksheet, minimum_columns=4)
            doc.values = [
                [worksheet.cell(row, col).value for col in range(1, last_col + 1)]
                for row in range(1, last_row + 1)
            ]
            workbook.close()
        except Exception as exc:  # openpyxl 的损坏文件异常类型较多
            doc.load_error = f"无法读取工作簿：{exc}"
            doc.issues = [CellIssue(0, 0, doc.load_error, "format")]
        return doc

    @staticmethod
    def _find_data_sheet(workbook: Any) -> tuple[Any, int, dict[str, int], str]:
        best: tuple[Any, int, dict[str, int]] | None = None
        best_score = -1
        for worksheet in workbook.worksheets:
            for row in range(1, min(worksheet.max_row, 20) + 1):
                found: dict[str, int] = {}
                supplemental: dict[str, int] = {}
                for col in range(1, min(worksheet.max_column, 40) + 1):
                    label = normalize_text(worksheet.cell(row, col).value)
                    if label in REQUIRED_COLUMNS:
                        found[label] = col
                    if label == WEEKEND_MONITOR_COLUMN:
                        found[WEEKEND_MONITOR_COLUMN] = col
                    if label == "姓名":
                        supplemental["姓名"] = col
                    elif _is_holiday_source_header(label):
                        supplemental[HOLIDAY_COLUMN] = col
                    elif _is_weekend_source_header(label):
                        supplemental[WEEKEND_COLUMN] = col
                if len(supplemental) == len(SUPPLEMENTAL_COLUMNS):
                    return worksheet, row, supplemental, SUPPLEMENTAL_DOCUMENT
                if len(found) > best_score:
                    best = worksheet, row, found
                    best_score = len(found)
                if all(label in found for label in REQUIRED_COLUMNS):
                    return worksheet, row, found, CLASS_DOCUMENT
        if best is None:
            worksheet = workbook.active
            return worksheet, 1, {}, CLASS_DOCUMENT
        worksheet, row, found = best
        return worksheet, row, found, CLASS_DOCUMENT

    @staticmethod
    def _used_bounds(worksheet: Any, minimum_columns: int) -> tuple[int, int]:
        last_row = 1
        last_col = minimum_columns
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    last_row = max(last_row, cell.row)
                    last_col = max(last_col, cell.column)
        return last_row, last_col


class WorkbookValidator:
    def __init__(self, template_names: Iterable[str]):
        self.template_names = {normalize_text(name) for name in template_names if normalize_text(name)}

    @classmethod
    def from_template(cls, template_path: Path) -> "WorkbookValidator":
        workbook = load_workbook(template_path, data_only=True, read_only=True)
        worksheet = workbook.active
        name_col = None
        header_row = 1
        for row in range(1, min(worksheet.max_row, 20) + 1):
            for col in range(1, min(worksheet.max_column, 40) + 1):
                if normalize_text(worksheet.cell(row, col).value) == "姓名":
                    name_col = col
                    header_row = row
                    break
            if name_col:
                break
        if name_col is None:
            workbook.close()
            raise ValueError("模板中没有找到“姓名”列")
        names = [
            worksheet.cell(row, name_col).value
            for row in range(header_row + 1, worksheet.max_row + 1)
        ]
        workbook.close()
        normalized = [normalize_text(name) for name in names if normalize_text(name)]
        duplicates = sorted({name for name in normalized if normalized.count(name) > 1})
        if duplicates:
            raise ValueError(f"模板姓名重复：{'、'.join(duplicates)}")
        return cls(normalized)

    def validate(self, doc: WorkbookDocument, expected: ExpectedTotals) -> list[CellIssue]:
        if doc.load_error:
            return doc.issues

        self._refresh_structure(doc)

        if doc.document_type == SUPPLEMENTAL_DOCUMENT:
            return self._validate_supplemental(doc)

        issues: list[CellIssue] = []
        missing = [name for name in REQUIRED_COLUMNS if name not in doc.columns]
        if missing:
            for name in missing:
                issues.append(
                    CellIssue(
                        doc.header_row,
                        self._first_blank_header_column(doc),
                        f"缺少必需列“{name}”",
                        "format",
                    )
                )
            doc.issues = issues
            return issues

        totals = {name: 0.0 for name in expected.as_dict()}
        for row in range(doc.header_row + 1, len(doc.values) + 1):
            name_col = doc.columns["姓名"]
            name_value = doc.cell(row, name_col)
            metric_labels = list(REQUIRED_COLUMNS[1:])
            if WEEKEND_MONITOR_COLUMN in doc.columns:
                metric_labels.append(WEEKEND_MONITOR_COLUMN)
            metric_values = [doc.cell(row, doc.columns[label]) for label in metric_labels]
            if normalize_text(name_value) == "" and all(
                value is None or (isinstance(value, str) and not value.strip()) for value in metric_values
            ):
                continue
            if not normalize_text(name_value):
                issues.append(CellIssue(row, name_col, "姓名不能为空", "name"))
            elif normalize_text(name_value) not in self.template_names:
                recommendation = self.recommend_name(name_value)
                recommendation_text = ""
                if recommendation:
                    suggested_name, similarity = recommendation
                    recommendation_text = f"；推荐“{suggested_name}”（相似度 {similarity:.0%}）"
                issues.append(
                    CellIssue(
                        row,
                        name_col,
                        f"姓名“{name_value}”未在模板中找到{recommendation_text}",
                        "name",
                    )
                )

            for label, value in zip(metric_labels, metric_values):
                col = doc.columns[label]
                try:
                    if label in totals:
                        totals[label] += number_from_cell(value)
                    else:
                        number_from_cell(value)
                except (TypeError, ValueError) as exc:
                    issues.append(CellIssue(row, col, f"{label}数据无效：{exc}", "number"))

        for label, expected_total in expected.as_dict().items():
            if label not in doc.columns:
                continue
            if not math.isclose(totals[label], expected_total, rel_tol=0, abs_tol=1e-9):
                issues.append(
                    CellIssue(
                        doc.header_row,
                        doc.columns[label],
                        f"{label}合计应为 {display_number(expected_total)}，实际为 {display_number(totals[label])}",
                        "total",
                    )
                )
        doc.issues = issues
        return issues

    def _validate_supplemental(self, doc: WorkbookDocument) -> list[CellIssue]:
        issues: list[CellIssue] = []
        missing = [name for name in SUPPLEMENTAL_COLUMNS if name not in doc.columns]
        if missing:
            for name in missing:
                issues.append(
                    CellIssue(
                        doc.header_row,
                        self._first_blank_header_column(doc),
                        f"缺少必需列“{name}”",
                        "format",
                    )
                )
            doc.issues = issues
            return issues

        seen_names: dict[str, int] = {}
        for row in range(doc.header_row + 1, len(doc.values) + 1):
            name_col = doc.columns["姓名"]
            name_value = doc.cell(row, name_col)
            metric_values = [doc.cell(row, doc.columns[label]) for label in SUPPLEMENTAL_COLUMNS[1:]]
            if normalize_text(name_value) == "" and all(
                value is None or (isinstance(value, str) and not value.strip())
                for value in metric_values
            ):
                continue

            name = normalize_text(name_value)
            if not name:
                issues.append(CellIssue(row, name_col, "姓名不能为空", "name"))
            elif name not in self.template_names:
                recommendation = self.recommend_name(name_value)
                recommendation_text = ""
                if recommendation:
                    suggested_name, similarity = recommendation
                    recommendation_text = f"；推荐“{suggested_name}”（相似度 {similarity:.0%}）"
                issues.append(
                    CellIssue(
                        row,
                        name_col,
                        f"姓名“{name_value}”未在模板中找到{recommendation_text}",
                        "name",
                    )
                )
            elif name in seen_names:
                issues.append(
                    CellIssue(
                        row,
                        name_col,
                        f"姓名“{name_value}”与第 {seen_names[name]} 行重复",
                        "name",
                    )
                )
            else:
                seen_names[name] = row

            for label, value in zip(SUPPLEMENTAL_COLUMNS[1:], metric_values):
                col = doc.columns[label]
                try:
                    number_from_cell(value)
                except (TypeError, ValueError) as exc:
                    issues.append(CellIssue(row, col, f"{label}数据无效：{exc}", "number"))

        doc.issues = issues
        return issues

    def recommend_name(self, name: Any) -> tuple[str, float] | None:
        return recommend_template_name(name, self.template_names)

    @staticmethod
    def _refresh_structure(doc: WorkbookDocument) -> None:
        best_row = doc.header_row
        best_columns: dict[str, int] = {}
        for row in range(1, min(len(doc.values), 20) + 1):
            columns: dict[str, int] = {}
            for col, value in enumerate(doc.values[row - 1], start=1):
                label = normalize_text(value)
                if doc.document_type == SUPPLEMENTAL_DOCUMENT:
                    if label == "姓名":
                        columns["姓名"] = col
                    elif _is_holiday_source_header(label):
                        columns[HOLIDAY_COLUMN] = col
                    elif _is_weekend_source_header(label):
                        columns[WEEKEND_COLUMN] = col
                elif label in REQUIRED_COLUMNS or label == WEEKEND_MONITOR_COLUMN:
                    columns[label] = col
            if len(columns) > len(best_columns):
                best_row = row
                best_columns = columns
            required = (
                SUPPLEMENTAL_COLUMNS
                if doc.document_type == SUPPLEMENTAL_DOCUMENT
                else REQUIRED_COLUMNS
            )
            if all(label in columns for label in required):
                break
        doc.header_row = best_row
        doc.columns = best_columns

    @staticmethod
    def _first_blank_header_column(doc: WorkbookDocument) -> int:
        width = max((len(row) for row in doc.values), default=4)
        for col in range(1, width + 1):
            if normalize_text(doc.cell(doc.header_row, col)) == "":
                return col
        # 没有空白表头时标记 A 列，确保问题在预览范围内可见。
        return 1


class TemplateRepository:
    def __init__(self, template_path: Path):
        self.template_path = template_path

    @classmethod
    def install_template(cls, source: Path, destination: Path) -> Path:
        """校验用户选择的工资模板，再复制为固定的 module.xlsx。"""
        source = Path(source)
        destination = Path(destination)
        if source.suffix.lower() != ".xlsx":
            raise ValueError("首次导入的模板必须是 .xlsx 文件")
        if destination.exists():
            raise FileExistsError(f"模板已存在：{destination}")

        repository = cls(source)
        repository.validate_structure()
        WorkbookValidator.from_template(source)

        destination.parent.mkdir(parents=True, exist_ok=True)
        temporary = destination.with_name(
            f".{destination.stem}-{uuid.uuid4().hex}.tmp{destination.suffix}"
        )
        try:
            shutil.copy2(source, temporary)
            make_writable(temporary)
            os.replace(temporary, destination)
            make_writable(destination)
        finally:
            if temporary.exists():
                temporary.unlink()
        return destination

    @classmethod
    def replace_template(
        cls,
        source: Path,
        destination: Path,
        rates: SalaryRates | None = None,
    ) -> Path:
        """校验候选模板后原子替换现有模板，失败时保留原文件。"""
        source = Path(source)
        destination = Path(destination)
        if source.suffix.lower() != ".xlsx":
            raise ValueError("工资模板必须是 .xlsx 文件")

        repository = cls(source)
        repository.validate_structure()
        WorkbookValidator.from_template(source)

        destination.parent.mkdir(parents=True, exist_ok=True)
        temporary = destination.with_name(
            f".{destination.stem}-{uuid.uuid4().hex}.tmp{destination.suffix}"
        )
        try:
            shutil.copy2(source, temporary)
            make_writable(temporary)
            if rates is not None:
                cls(temporary).update_rate_headers(rates)
            os.replace(temporary, destination)
            make_writable(destination)
        finally:
            if temporary.exists():
                temporary.unlink()
        return destination

    def validate_structure(self) -> None:
        workbook = load_workbook(self.template_path, data_only=False, read_only=True)
        try:
            worksheet = self._find_template_sheet(workbook)
            labels = [label for _, label in self._headers(worksheet) if label]
            required = {
                "姓名",
                "G1晚",
                "G2晚",
                "G3晚",
                "晚读",
                "高中中午",
                "全程下午",
                WEEKEND_COLUMN,
                HOLIDAY_COLUMN,
                "合计金额",
            }
            missing = sorted(required - set(labels))
            amount_prefixes = (
                "G1晚金额",
                "G2晚金额",
                "G3晚金额",
                "晚读金额",
                "中午金额",
                "下午金额",
            )
            missing.extend(
                prefix for prefix in amount_prefixes if not any(label.startswith(prefix) for label in labels)
            )
            if missing:
                raise ValueError(f"所选文件不是完整工资模板，缺少列：{'、'.join(missing)}")
        finally:
            workbook.close()

    def update_rate_headers(self, rates: SalaryRates) -> None:
        workbook = load_workbook(self.template_path)
        try:
            worksheet = self._find_template_sheet(workbook)
            headers = self._headers(worksheet)
            replacements = {
                "G1晚金额": rates.grade_evening["1"],
                "G2晚金额": rates.grade_evening["2"],
                "G3晚金额": rates.grade_evening["3"],
                "晚读金额": rates.evening_reading,
                "中午金额": rates.noon,
                "下午金额": rates.afternoon,
            }
            found = set()
            for cell, label in headers:
                for prefix, rate in replacements.items():
                    if label.startswith(prefix):
                        cell.value = f"{prefix}{_rate_label(rate)}"
                        found.add(prefix)
            missing = set(replacements) - found
            if missing:
                raise ValueError(f"模板缺少金额列：{'、'.join(sorted(missing))}")
            self._save_atomic(workbook)
        finally:
            workbook.close()

    def append_person(self, name: Any) -> int:
        """把新姓名追加到模板数据区末尾，并返回新行号。"""
        normalized_name = normalize_text(name)
        if not normalized_name:
            raise ValueError("追加到模板的姓名不能为空")

        workbook = load_workbook(self.template_path)
        try:
            worksheet = self._find_template_sheet(workbook)
            header_cells = self._headers(worksheet)
            columns = {label: cell.column for cell, label in header_cells if label}
            if "姓名" not in columns:
                raise ValueError("模板中没有找到“姓名”列")
            header_row = header_cells[0][0].row if header_cells else 1
            name_column = columns["姓名"]

            existing_names: dict[str, int] = {}
            for row in range(header_row + 1, worksheet.max_row + 1):
                existing_name = normalize_text(worksheet.cell(row, name_column).value)
                if existing_name:
                    existing_names[existing_name] = row
            if normalized_name in existing_names:
                raise ValueError(
                    f"姓名“{normalized_name}”已在模板第 {existing_names[normalized_name]} 行"
                )

            last_row = max(existing_names.values(), default=header_row)
            new_row = last_row + 1
            if last_row > header_row:
                self._copy_row_format(worksheet, last_row, new_row)

            for column in range(1, worksheet.max_column + 1):
                worksheet.cell(new_row, column).value = 0
            worksheet.cell(new_row, name_column).value = normalized_name

            sequence_column = columns.get("序号")
            if sequence_column is not None:
                sequence_values = []
                for row in range(header_row + 1, last_row + 1):
                    value = worksheet.cell(row, sequence_column).value
                    if isinstance(value, (int, float)) and not isinstance(value, bool):
                        sequence_values.append(int(value))
                worksheet.cell(new_row, sequence_column).value = (
                    max(sequence_values, default=0) + 1
                )

            self._save_atomic(workbook)
            return new_row
        finally:
            workbook.close()

    def _save_atomic(self, workbook: Any) -> None:
        make_writable(self.template_path)
        temporary = self.template_path.with_name(
            f".{self.template_path.stem}-{uuid.uuid4().hex}.tmp{self.template_path.suffix}"
        )
        try:
            workbook.save(temporary)
            make_writable(temporary)
            os.replace(temporary, self.template_path)
        finally:
            if temporary.exists():
                temporary.unlink()

    @staticmethod
    def _copy_row_format(worksheet: Any, source_row: int, target_row: int) -> None:
        source_dimension = worksheet.row_dimensions[source_row]
        target_dimension = worksheet.row_dimensions[target_row]
        target_dimension.height = source_dimension.height
        target_dimension.hidden = source_dimension.hidden
        target_dimension.outlineLevel = source_dimension.outlineLevel
        for column in range(1, worksheet.max_column + 1):
            source = worksheet.cell(source_row, column)
            target = worksheet.cell(target_row, column)
            if source.has_style:
                target.font = copy(source.font)
                target.fill = copy(source.fill)
                target.border = copy(source.border)
                target.alignment = copy(source.alignment)
                target.number_format = source.number_format
                target.protection = copy(source.protection)

    @staticmethod
    def _find_template_sheet(workbook: Any) -> Any:
        for worksheet in workbook.worksheets:
            for row in range(1, min(worksheet.max_row, 20) + 1):
                labels = {normalize_text(worksheet.cell(row, col).value) for col in range(1, worksheet.max_column + 1)}
                if {"姓名", "合计金额"}.issubset(labels):
                    return worksheet
        raise ValueError("模板中没有找到工资表")

    @staticmethod
    def _headers(worksheet: Any) -> list[tuple[Any, str]]:
        for row in range(1, min(worksheet.max_row, 20) + 1):
            labels = [normalize_text(worksheet.cell(row, col).value) for col in range(1, worksheet.max_column + 1)]
            if "姓名" in labels:
                return [(worksheet.cell(row, col), labels[col - 1]) for col in range(1, worksheet.max_column + 1)]
        return []


def _rate_label(rate: float) -> str:
    return str(int(rate)) if float(rate).is_integer() else str(rate)


class SalaryCalculator:
    def __init__(self, template_path: Path, rates: SalaryRates):
        self.template_path = template_path
        self.rates = rates

    def calculate(
        self,
        documents: Iterable[WorkbookDocument],
        grade: str,
        output_path: Path,
    ) -> CalculationResult:
        if grade not in {"1", "2", "3"}:
            raise ValueError("年级必须是 1、2 或 3")
        documents = list(documents)
        class_documents = [doc for doc in documents if doc.document_type == CLASS_DOCUMENT]
        supplemental_documents = [
            doc for doc in documents if doc.document_type == SUPPLEMENTAL_DOCUMENT
        ]
        if not class_documents:
            raise ValueError("没有可计算的班级文件")
        if len(supplemental_documents) > 1:
            raise ValueError("只能导入一份周六日及节假日年级统计表")
        if any(doc.issues for doc in documents):
            raise ValueError("仍有未解决的数据问题，不能生成工资表")

        shutil.copy2(self.template_path, output_path)
        make_writable(output_path)
        workbook = load_workbook(output_path)
        worksheet = TemplateRepository._find_template_sheet(workbook)
        header_row, columns = self._template_columns(worksheet)
        self._apply_rate_headers(worksheet, header_row, columns)

        grade_count = f"G{grade}晚"
        grade_amount = self._amount_column(columns, f"G{grade}晚金额")
        evening_amount = self._amount_column(columns, "晚读金额")
        noon_amount = self._amount_column(columns, "中午金额")
        afternoon_amount = self._amount_column(columns, "下午金额")
        required = [
            "姓名",
            grade_count,
            "晚读",
            "高中中午",
            "全程下午",
            WEEKEND_COLUMN,
            HOLIDAY_COLUMN,
            "合计金额",
        ]
        missing = [label for label in required if label not in columns]
        if missing:
            workbook.close()
            raise ValueError(f"模板缺少列：{'、'.join(missing)}")

        aggregates: dict[str, dict[str, float]] = defaultdict(
            lambda: {
                "晚自修": 0.0,
                "高中中午": 0.0,
                "全程下午": 0.0,
                WEEKEND_COLUMN: 0.0,
                HOLIDAY_COLUMN: 0.0,
            }
        )
        for doc in class_documents:
            for row in range(doc.header_row + 1, len(doc.values) + 1):
                name = normalize_text(doc.cell(row, doc.columns["姓名"]))
                if not name:
                    continue
                for label in REQUIRED_COLUMNS[1:]:
                    aggregates[name][label] += number_from_cell(doc.cell(row, doc.columns[label]))
                if WEEKEND_MONITOR_COLUMN in doc.columns:
                    aggregates[name][WEEKEND_COLUMN] += number_from_cell(
                        doc.cell(row, doc.columns[WEEKEND_MONITOR_COLUMN])
                    )

        if supplemental_documents:
            supplemental = supplemental_documents[0]
            for row in range(supplemental.header_row + 1, len(supplemental.values) + 1):
                name = normalize_text(supplemental.cell(row, supplemental.columns["姓名"]))
                if not name:
                    continue
                aggregates[name][WEEKEND_COLUMN] += number_from_cell(
                    supplemental.cell(row, supplemental.columns[WEEKEND_COLUMN])
                )
                aggregates[name][HOLIDAY_COLUMN] = number_from_cell(
                    supplemental.cell(row, supplemental.columns[HOLIDAY_COLUMN])
                )

        teacher_rows: dict[str, int] = {}
        for row in range(header_row + 1, worksheet.max_row + 1):
            name = normalize_text(worksheet.cell(row, columns["姓名"]).value)
            if name:
                teacher_rows[name] = row

        count_columns = [
            label
            for label in (
                "G1晚",
                "G2晚",
                "G3晚",
                "晚读",
                "高中中午",
                "全程下午",
                WEEKEND_COLUMN,
                HOLIDAY_COLUMN,
            )
            if label in columns
        ]
        amount_columns = [
            label
            for label in columns
            if label.startswith(("G1晚金额", "G2晚金额", "G3晚金额", "晚读金额", "中午金额", "下午金额"))
        ] + ["合计金额"]
        for row in teacher_rows.values():
            for label in count_columns + amount_columns:
                worksheet.cell(row, columns[label]).value = 0

        for name, values in aggregates.items():
            row = teacher_rows[name]
            evening = values["晚自修"]
            noon = values["高中中午"]
            afternoon = values["全程下午"]
            worksheet.cell(row, columns[grade_count]).value = display_number(evening)
            worksheet.cell(row, columns["晚读"]).value = display_number(evening)
            worksheet.cell(row, columns["高中中午"]).value = display_number(noon)
            worksheet.cell(row, columns["全程下午"]).value = display_number(afternoon)
            worksheet.cell(row, columns[WEEKEND_COLUMN]).value = display_number(
                values[WEEKEND_COLUMN] * self.rates.weekend
            )
            worksheet.cell(row, columns[HOLIDAY_COLUMN]).value = display_number(
                values[HOLIDAY_COLUMN] * self.rates.holiday
            )
            worksheet.cell(row, columns[grade_amount]).value = display_number(
                evening * self.rates.grade_evening[grade]
            )
            worksheet.cell(row, columns[evening_amount]).value = display_number(
                evening * self.rates.evening_reading
            )
            worksheet.cell(row, columns[noon_amount]).value = display_number(noon * self.rates.noon)
            worksheet.cell(row, columns[afternoon_amount]).value = display_number(
                afternoon * self.rates.afternoon
            )

        all_amount_columns = [
            label
            for label in columns
            if label.startswith(("G1晚金额", "G2晚金额", "G3晚金额", "晚读金额", "中午金额", "下午金额"))
        ] + [WEEKEND_COLUMN, HOLIDAY_COLUMN]
        total_amount = 0.0
        for row in teacher_rows.values():
            total = sum(float(worksheet.cell(row, columns[label]).value or 0) for label in all_amount_columns)
            worksheet.cell(row, columns["合计金额"]).value = display_number(total)
            total_amount += total

        self._style_output(worksheet, header_row, columns, teacher_rows.values())
        workbook.save(output_path)
        workbook.close()

        values = self._read_preview(output_path)
        return CalculationResult(
            workbook_path=output_path,
            sheet_name=worksheet.title,
            values=values,
            total_amount=total_amount,
            teacher_count=sum(1 for value in aggregates.values() if any(value.values())),
            source_file_count=len(class_documents),
            supplemental_file_count=len(supplemental_documents),
        )

    @staticmethod
    def _template_columns(worksheet: Any) -> tuple[int, dict[str, int]]:
        for row in range(1, min(worksheet.max_row, 20) + 1):
            columns: dict[str, int] = {}
            for col in range(1, worksheet.max_column + 1):
                label = normalize_text(worksheet.cell(row, col).value)
                if label:
                    columns[label] = col
            if "姓名" in columns and "合计金额" in columns:
                return row, columns
        raise ValueError("模板表头无效")

    @staticmethod
    def _amount_column(columns: dict[str, int], prefix: str) -> str:
        matches = [label for label in columns if label.startswith(prefix)]
        if not matches:
            raise ValueError(f"模板缺少“{prefix}”列")
        return matches[0]

    def _apply_rate_headers(
        self, worksheet: Any, header_row: int, columns: dict[str, int]
    ) -> None:
        rates = {
            "G1晚金额": self.rates.grade_evening["1"],
            "G2晚金额": self.rates.grade_evening["2"],
            "G3晚金额": self.rates.grade_evening["3"],
            "晚读金额": self.rates.evening_reading,
            "中午金额": self.rates.noon,
            "下午金额": self.rates.afternoon,
        }
        for prefix, rate in rates.items():
            label = self._amount_column(columns, prefix)
            col = columns.pop(label)
            new_label = f"{prefix}{_rate_label(rate)}"
            worksheet.cell(header_row, col).value = new_label
            columns[new_label] = col

    @staticmethod
    def _style_output(
        worksheet: Any, header_row: int, columns: dict[str, int], rows: Iterable[int]
    ) -> None:
        header_fill = PatternFill("solid", fgColor="1F4E78")
        for col in range(1, worksheet.max_column + 1):
            cell = worksheet.cell(header_row, col)
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        worksheet.freeze_panes = worksheet.cell(header_row + 1, 3)
        for row in rows:
            for col in range(1, worksheet.max_column + 1):
                worksheet.cell(row, col).alignment = Alignment(horizontal="center", vertical="center")
            for label, col in columns.items():
                if "金额" in label or label in {WEEKEND_COLUMN, HOLIDAY_COLUMN}:
                    worksheet.cell(row, col).number_format = '#,##0.00'
        worksheet.column_dimensions["B"].width = max(12, worksheet.column_dimensions["B"].width or 0)

    @staticmethod
    def _read_preview(path: Path) -> list[list[Any]]:
        workbook = load_workbook(path, data_only=True, read_only=True)
        worksheet = TemplateRepository._find_template_sheet(workbook)
        values = [list(row) for row in worksheet.iter_rows(values_only=True)]
        workbook.close()
        return values


def save_document_edits(doc: WorkbookDocument) -> None:
    if not doc.cache_path or doc.load_error:
        return
    make_writable(doc.cache_path)
    temporary = doc.cache_path.with_name(f".{doc.cache_path.name}.tmp")
    workbook = load_workbook(doc.cache_path)
    try:
        worksheet = workbook[doc.sheet_name]
        for row_index, row in enumerate(doc.values, start=1):
            for col_index, value in enumerate(row, start=1):
                worksheet.cell(row_index, col_index).value = value
        workbook.save(temporary)
    finally:
        workbook.close()
    make_writable(temporary)
    os.replace(temporary, doc.cache_path)
    try_save_document_to_source(doc)


def try_save_document_to_source(doc: WorkbookDocument) -> bool:
    """尝试把缓存中的修改同步回导入路径，失败时静默保留缓存副本。"""
    if not doc.cache_path or doc.load_error or doc.cache_path == doc.source_path:
        return False

    source = doc.source_path
    temporary: Path | None = None
    original_mode: int | None = None
    replaced = False
    try:
        temporary = source.with_name(
            f".{source.stem}-{uuid.uuid4().hex}.tmp{source.suffix}"
        )
        shutil.copy2(doc.cache_path, temporary)
        if source.exists():
            original_mode = source.stat().st_mode
            if not original_mode & stat.S_IWRITE:
                make_writable(source)
        os.replace(temporary, source)
        replaced = True
        if original_mode is not None:
            try:
                source.chmod(original_mode)
            except Exception:
                pass
        return True
    except Exception:
        if not replaced and original_mode is not None:
            try:
                if source.exists():
                    source.chmod(original_mode)
            except Exception:
                pass
        return False
    finally:
        try:
            if temporary is not None and temporary.exists():
                temporary.unlink()
        except Exception:
            pass


def discover_excel_files(folder: Path) -> list[Path]:
    files = [
        path
        for path in folder.iterdir()
        if path.is_file()
        and path.suffix.lower() in ALL_EXCEL_SUFFIXES
        and not path.name.startswith("~$")
    ]
    return sorted(files, key=_natural_sort_key)


def find_name_occurrences(
    documents: Iterable[WorkbookDocument], name: Any
) -> list[tuple[WorkbookDocument, int, int]]:
    """查找所有导入文档中与指定姓名完全相同（忽略空白）的单元格。"""
    target = normalize_text(name)
    if not target:
        return []
    matches: list[tuple[WorkbookDocument, int, int]] = []
    for document in documents:
        if document.load_error:
            continue
        name_column = document.columns.get("姓名")
        if name_column is None:
            continue
        for row in range(document.header_row + 1, len(document.values) + 1):
            if normalize_text(document.cell(row, name_column)) == target:
                matches.append((document, row, name_column))
    return matches


def _natural_sort_key(path: Path) -> list[Any]:
    return [int(part) if part.isdigit() else part.lower() for part in re.split(r"(\d+)", path.name)]
